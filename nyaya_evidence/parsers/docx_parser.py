"""DOCX document parser for NYAYA-SATYA using python-docx."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import docx

from nyaya_evidence.contracts.evidence import MediaType
from nyaya_evidence.parsers.base import (
    BaseDocumentParser,
    DocumentParserError,
    ExtractionMetadata,
    ParsedDocument,
    ParsedPage,
)


class DocxDocumentParser(BaseDocumentParser):
    @property
    def supported_media_types(self) -> set[MediaType]:
        return {MediaType.APPLICATION_DOCX}

    @property
    def parser_name(self) -> str:
        return "DocxDocumentParser"

    @property
    def parser_version(self) -> str:
        return "python-docx-parser@1.0.0"

    def parse(self, *, evidence_id: str, raw_bytes: bytes, source_hash: str) -> ParsedDocument:
        stream = io.BytesIO(raw_bytes)
        try:
            doc = docx.Document(stream)
            elements: list[str] = []

            # Extract paragraphs
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    elements.append(text)

            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        elements.append(" | ".join(row_text))

            full_text = "\n\n".join(elements)
            extraction_hash = self.compute_extraction_hash(full_text)
            page = ParsedPage(page_number=1, text=full_text)

            meta = ExtractionMetadata(
                extraction_id=f"ext_{uuid.uuid4().hex[:16]}",
                evidence_id=evidence_id,
                media_type=MediaType.APPLICATION_DOCX,
                parser_name=self.parser_name,
                parser_version=self.parser_version,
                source_hash=source_hash,
                extraction_hash=extraction_hash,
                total_pages=1,
                char_count=len(full_text),
                extracted_at=datetime.now(UTC),
            )

            return ParsedDocument(evidence_id=evidence_id, text_content=full_text, pages=[page], metadata=meta)

        except Exception as exc:
            raise DocumentParserError(f"Failed to parse DOCX document: {exc}") from exc


__all__ = ["DocxDocumentParser"]
